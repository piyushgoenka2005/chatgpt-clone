from json import dumps, loads
from time import time
from flask import request, session, jsonify
from hashlib import sha256
from datetime import datetime
from requests import get, post
import os
import io
import base64
from pypdf import PdfReader
from docx import Document as DocxDocument

from server.config import special_instructions
from server.db import Database


def extract_document_text(doc_data: dict) -> str:
    """Extract plain text from uploaded documents (PDF, DOCX, TXT, CSV, JSON, code, etc.)."""
    if not doc_data:
        return ""
    
    filename = doc_data.get("name", "document")
    content = doc_data.get("content", "")
    ext = filename.lower().split(".")[-1] if "." in filename else ""
    
    if not isinstance(content, str):
        return ""

    if not content.startswith("data:"):
        return content
    
    try:
        header, encoded = content.split(",", 1)
        file_bytes = base64.b64decode(encoded)
        
        if ext == "pdf":
            pdf_stream = io.BytesIO(file_bytes)
            reader = PdfReader(pdf_stream)
            text_parts = []
            for page_idx, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(f"--- Page {page_idx + 1} ---\n{page_text}")
            return "\n\n".join(text_parts) if text_parts else "[Empty PDF]"
            
        elif ext in ["docx", "doc"]:
            docx_stream = io.BytesIO(file_bytes)
            doc = DocxDocument(docx_stream)
            text_parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    text_parts.append(" | ".join([cell.text.strip() for cell in row.cells]))
            return "\n".join(text_parts) if text_parts else "[Empty DOCX]"
            
        else:
            return file_bytes.decode("utf-8", errors="replace")
    except Exception as e:
        print(f"Error extracting document text for {filename}: {e}")
        return f"[Document text could not be extracted: {e}]"


class Backend_Api:
    def __init__(self, app, config: dict) -> None:
        self.app = app
        self.openai_key = os.getenv("OPENAI_API_KEY") or config.get('openai_key', '')
        self.openai_api_base = os.getenv("OPENAI_API_BASE") or config.get('openai_api_base', 'https://api.openai.com')
        self.proxy = config.get('proxy', {'enable': False})
        
        mongo_uri = os.getenv("MONGO_URI") or config.get('mongo_uri', '')
        mongo_db = os.getenv("MONGO_DB") or config.get('mongo_db', 'chatgpt_clone')
        try:
            self.db = Database(mongo_uri, mongo_db)
        except Exception as e:
            print("Notice on Mongo connection:", e)
            self.db = None

        self.routes = {
            '/backend-api/v2/conversation': {
                'function': self._conversation,
                'methods': ['POST']
            },
            '/backend-api/v2/auth/register': {
                'function': self._auth_register,
                'methods': ['POST']
            },
            '/backend-api/v2/auth/login': {
                'function': self._auth_login,
                'methods': ['POST']
            },
            '/backend-api/v2/auth/logout': {
                'function': self._auth_logout,
                'methods': ['POST']
            },
            '/backend-api/v2/auth/me': {
                'function': self._auth_me,
                'methods': ['GET']
            },
            '/backend-api/v2/conversations': {
                'function': self._get_conversations,
                'methods': ['GET']
            },
            '/backend-api/v2/conversation/<conversation_id>': {
                'function': self._handle_conversation_by_id,
                'methods': ['GET', 'DELETE']
            },
            '/backend-api/v2/conversation/<conversation_id>/message/<int:message_index>': {
                'function': self._delete_message,
                'methods': ['DELETE']
            },
            '/backend-api/v2/conversations/clear': {
                'function': self._clear_all_conversations,
                'methods': ['DELETE']
            }
        }

    def _get_current_user_id(self):
        return session.get('user_id')

    def _auth_register(self):
        data = request.json or {}
        name = data.get('name', '').strip()
        email = data.get('email', '').strip()
        password = data.get('password', '')

        success, res = self.db.register_user(name, email, password)
        if not success:
            return jsonify({'success': False, 'message': res}), 400

        session['user_id'] = res['user_id']
        session['user_email'] = res['email']
        session['user_name'] = res['name']
        return jsonify({'success': True, 'user': {'email': res['email'], 'name': res['name']}})

    def _auth_login(self):
        data = request.json or {}
        email = data.get('email', '').strip()
        password = data.get('password', '')

        success, res = self.db.authenticate_user(email, password)
        if not success:
            return jsonify({'success': False, 'message': res}), 401

        session['user_id'] = res['user_id']
        session['user_email'] = res['email']
        session['user_name'] = res.get('name', res['email'].split('@')[0])
        return jsonify({'success': True, 'user': {'email': res['email'], 'name': session['user_name']}})

    def _auth_logout(self):
        session.clear()
        return jsonify({'success': True})

    def _auth_me(self):
        user_id = self._get_current_user_id()
        if not user_id:
            return jsonify({'authenticated': False})
        user = self.db.get_user_by_id(user_id)
        if not user:
            session.clear()
            return jsonify({'authenticated': False})
        return jsonify({'authenticated': True, 'user': {'email': user['email'], 'name': user.get('name', '')}})

    def _get_conversations(self):
        user_id = self._get_current_user_id()
        if not user_id:
            return jsonify({'success': False, 'message': 'Unauthorized'}), 401
        
        conversations = self.db.get_conversations(user_id)
        return jsonify({'success': True, 'conversations': conversations})

    def _handle_conversation_by_id(self, conversation_id):
        user_id = self._get_current_user_id()
        if not user_id:
            return jsonify({'success': False, 'message': 'Unauthorized'}), 401

        if request.method == 'GET':
            doc = self.db.get_conversation(user_id, conversation_id)
            if not doc:
                return jsonify({'success': False, 'items': []})
            return jsonify({'success': True, 'conversation': doc, 'items': doc.get('items', [])})

        elif request.method == 'DELETE':
            success = self.db.delete_conversation(user_id, conversation_id)
            return jsonify({'success': success})

    def _delete_message(self, conversation_id, message_index):
        user_id = self._get_current_user_id()
        if not user_id:
            return jsonify({'success': False, 'message': 'Unauthorized'}), 401

        success = self.db.delete_message(user_id, conversation_id, message_index)
        return jsonify({'success': success})

    def _clear_all_conversations(self):
        user_id = self._get_current_user_id()
        if not user_id:
            return jsonify({'success': False, 'message': 'Unauthorized'}), 401

        self.db.delete_all_conversations(user_id)
        return jsonify({'success': True})

    def _conversation(self):
        try:
            user_id = self._get_current_user_id()
            if not user_id:
                return jsonify({'success': False, 'error': 'Please sign in to start chatting.'}), 401

            req_json = request.json or {}
            conv_id = req_json.get('conversation_id')
            jailbreak = req_json.get('jailbreak', 'default')
            meta_content = req_json.get('meta', {}).get('content', {})
            internet_access = meta_content.get('internet_access', False)
            _conversation = meta_content.get('conversation', [])
            parts = meta_content.get('parts', [{'content': '', 'role': 'user'}])
            prompt = parts[0] if parts else {'content': '', 'role': 'user'}
            image = meta_content.get('image') or prompt.get('image') or req_json.get('image')
            document = meta_content.get('document') or prompt.get('document') or req_json.get('document')

            from server.db import IST
            current_date = datetime.now(IST).strftime("%Y-%m-%d %I:%M %p IST")
            system_message = f'You are ChatGPT, a large language model trained by OpenAI. Strictly follow user instructions. You can analyze images and examine attached documents (PDF, DOCX, text files, datasets, code). Current date and time: {current_date}'

            doc_prompt_context = ""
            doc_meta_for_db = None
            if document and isinstance(document, dict):
                doc_name = document.get('name', 'Attached Document')
                raw_doc_text = extract_document_text(document)
                safe_doc_text = raw_doc_text[:40000] + ("\n\n...[Document content truncated for length]..." if len(raw_doc_text) > 40000 else "")
                doc_prompt_context = f"\n\n[ATTACHED DOCUMENT: {doc_name}]\n--- Document Content Start ---\n{safe_doc_text}\n--- Document Content End ---\n"
                doc_meta_for_db = {
                    'name': doc_name,
                    'size': document.get('size', 0),
                    'type': document.get('type', ''),
                    'snippet': safe_doc_text[:250] if safe_doc_text else ''
                }

            user_text = prompt.get("content", "").strip()
            if not user_text:
                if image and document:
                    user_text = "What is in this image and attached document?"
                elif image:
                    user_text = "What is in this image?"
                elif document:
                    user_text = f"Please read, analyze, and summarize the attached document ({doc_meta_for_db['name']})."

            # Append user message to MongoDB with image & document metadata
            if user_id and conv_id and (user_text or image or document):
                try:
                    self.db.append_message(user_id, conv_id, "user", user_text, image=image, document=doc_meta_for_db)
                except Exception as db_err:
                    print("MongoDB user message append error:", db_err)

            # Format previous conversation for OpenAI (including previous images)
            formatted_conversation = []
            for item in _conversation:
                role = item.get('role', 'user')
                content_text = item.get('content', '')
                item_image = item.get('image')
                if item_image and role == 'user':
                    formatted_conversation.append({
                        'role': role,
                        'content': [
                            {'type': 'text', 'text': content_text or 'Image'},
                            {'type': 'image_url', 'image_url': {'url': item_image, 'detail': 'auto'}}
                        ]
                    })
                else:
                    formatted_conversation.append({
                        'role': role,
                        'content': content_text
                    })

            # Format current user message (combining text and document context)
            full_prompt_text = (user_text + doc_prompt_context).strip()
            if image:
                current_user_msg = {
                    'role': 'user',
                    'content': [
                        {'type': 'text', 'text': full_prompt_text},
                        {'type': 'image_url', 'image_url': {'url': image, 'detail': 'auto'}}
                    ]
                }
            else:
                current_user_msg = {
                    'role': 'user',
                    'content': full_prompt_text
                }

            extra = []
            if internet_access and user_text:
                try:
                    search = get('https://ddg-api.herokuapp.com/search', params={
                        'query': user_text,
                        'limit': 3,
                    }, timeout=5)

                    if search.status_code == 200:
                        results = search.json()
                        blob = ''
                        for index, result in enumerate(results):
                            blob += f'[{index}] "{result.get("snippet", "")}"\nURL:{result.get("link", "")}\n\n'

                        date = datetime.now().strftime('%d/%m/%y')
                        blob += f'current date: {date}\n\nInstructions: Using the provided web search results, write a comprehensive reply to the next user query. Make sure to cite results using [[number](URL)] notation after the reference. If the provided search results refer to multiple subjects with the same name, write separate answers for each subject. Ignore your previous response if any.'
                        extra = [{'role': 'user', 'content': blob}]
                except Exception as search_err:
                    print(f"Web search error (bypassed): {search_err}")

            special_prompt = special_instructions.get(jailbreak, [])
            conversation = [{'role': 'system', 'content': system_message}] + \
                extra + special_prompt + \
                formatted_conversation + [current_user_msg]

            url = f"{self.openai_api_base}/v1/chat/completions"

            proxies = None
            if isinstance(self.proxy, dict) and self.proxy.get('enable'):
                proxies = {
                    'http': self.proxy.get('http'),
                    'https': self.proxy.get('https'),
                }

            # Enforce gpt-4o-mini exclusively as requested
            model_to_use = 'gpt-4o-mini'

            gpt_resp = post(
                url     = url,
                proxies = proxies,
                headers = {
                    'Authorization': f'Bearer {self.openai_key}'
                }, 
                json    = {
                    'model'             : model_to_use, 
                    'messages'          : conversation,
                    'stream'            : True
                },
                stream  = True,
                timeout = 60
            )

            if gpt_resp.status_code >= 400:
                try:
                    error_data = gpt_resp.json().get('error', {})
                    error_message = error_data.get('message', "An error occurred with OpenAI API")
                except Exception:
                    error_message = gpt_resp.text or f"OpenAI error {gpt_resp.status_code}"
                return {
                    'success': False,
                    'error': error_message,
                    'status_code': gpt_resp.status_code
                }, gpt_resp.status_code

            def stream():
                full_reply = []
                for chunk in gpt_resp.iter_lines():
                    if not chunk:
                        continue
                    try:
                        decoded = chunk.decode("utf-8").strip()
                        if not decoded.startswith("data: "):
                            continue
                        data_str = decoded[6:].strip()
                        if data_str == "[DONE]":
                            break
                        decoded_line = loads(data_str)
                        choices = decoded_line.get("choices", [])
                        if choices:
                            token = choices[0].get("delta", {}).get("content")
                            if token is not None:
                                full_reply.append(token)
                                yield token
                    except GeneratorExit:
                        break
                    except Exception:
                        continue

                # Append complete assistant reply to MongoDB
                if user_id and conv_id and full_reply:
                    try:
                        self.db.append_message(user_id, conv_id, "assistant", "".join(full_reply))
                    except Exception as db_err:
                        print("MongoDB assistant message append error:", db_err)

            return self.app.response_class(stream(), mimetype='text/event-stream')

        except Exception as e:
            print(f"Conversation error: {e}")
            return {
                '_action': '_ask',
                'success': False,
                'error': f"an error occurred: {str(e)}"}, 400
